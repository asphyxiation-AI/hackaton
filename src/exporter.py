"""Экспорт дайджеста в Markdown и DOCX."""

import re
from pathlib import Path

import docx


def _add_runs(paragraph, line: str) -> None:
    for index, chunk in enumerate(line.split("**")):
        if chunk:
            run = paragraph.add_run(chunk)
            run.bold = index % 2 == 1


def export_to_markdown(text: str, file_path: str) -> str:
    """Сохраняет текст в .md файл и возвращает путь."""
    path = Path(file_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return str(path)


def export_to_docx(text: str, file_path: str) -> str:
    """Конвертирует Markdown-дайджест в .docx и возвращает путь."""
    document = docx.Document()
    for raw in text.replace("\r", "").splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        stripped = line.lstrip()
        if stripped.startswith("### "):
            paragraph = document.add_heading(level=3)
            stripped = stripped[4:]
        elif stripped.startswith("## "):
            paragraph = document.add_heading(level=2)
            stripped = stripped[3:]
        elif stripped.startswith("# "):
            paragraph = document.add_heading(level=1)
            stripped = stripped[2:]
        elif stripped.startswith(("- ", "* ")):
            paragraph = document.add_paragraph(style="List Bullet")
            stripped = stripped[2:]
        elif re.match(r"^\d+[.)]\s+", stripped):
            paragraph = document.add_paragraph(style="List Number")
            stripped = re.sub(r"^\d+[.)]\s+", "", stripped)
        else:
            paragraph = document.add_paragraph()
        _add_runs(paragraph, stripped)
    document.save(file_path)
    return str(file_path)